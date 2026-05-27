"""
reconpilot.mapping — parse Functional/Technical mapping documents.

Supports Excel (.xlsx) and CSV mapping specs used by BA and technical teams.

Standard mapping doc format (columns):
  source_table  | source_column | target_table | target_column | transformation | is_key | nullable | data_type | notes

Both column names and their common enterprise aliases are supported.
"""

from __future__ import annotations
import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional


# Column aliases — handles variations in naming conventions across teams
_COL_ALIASES = {
    "source_table":     ["source_table", "src_table", "source table", "from_table", "from table"],
    "source_column":    ["source_column", "src_column", "source_field", "src_field", "source field", "from_column"],
    "target_table":     ["target_table", "tgt_table", "target table", "to_table", "destination_table"],
    "target_column":    ["target_column", "tgt_column", "target_field", "tgt_field", "target field", "to_column"],
    "transformation":   ["transformation", "transform", "mapping_rule", "rule", "logic", "mapping logic"],
    "is_key":           ["is_key", "key", "primary_key", "pk", "is pk", "key_field"],
    "nullable":         ["nullable", "is_nullable", "optional", "null_allowed", "allow null"],
    "data_type":        ["data_type", "datatype", "type", "field_type", "tgt_datatype"],
    "notes":            ["notes", "comments", "remarks", "description"],
}


@dataclass
class FieldMapping:
    """A single field-level mapping from source to target."""
    source_table: str
    source_column: str
    target_table: str
    target_column: str
    transformation: str = "direct"   # direct | trim | upper | lower | date_format | custom
    is_key: bool = False
    nullable: bool = True
    data_type: str = ""
    notes: str = ""

    @property
    def is_direct(self) -> bool:
        return self.transformation.strip().lower() in ("direct", "none", "", "as-is", "as is")

    @property
    def has_transform(self) -> bool:
        return not self.is_direct


@dataclass
class MappingDocument:
    """Parsed mapping document — one or more tables, many field mappings."""
    source_path: str
    mappings: list[FieldMapping]
    warnings: list[str] = field(default_factory=list)

    @property
    def tables(self) -> list[str]:
        return sorted({m.source_table for m in self.mappings})

    @property
    def target_tables(self) -> list[str]:
        return sorted({m.target_table for m in self.mappings})

    def for_table(self, source_table: str) -> list[FieldMapping]:
        return [m for m in self.mappings if m.source_table.lower() == source_table.lower()]

    def keys_for_table(self, source_table: str) -> list[str]:
        return [m.source_column for m in self.for_table(source_table) if m.is_key]

    def target_keys_for_table(self, source_table: str) -> list[str]:
        return [m.target_column for m in self.for_table(source_table) if m.is_key]

    def summary(self):
        print(f"\n📋 Mapping Document: {self.source_path}")
        print(f"   Tables : {len(self.tables)}")
        print(f"   Fields : {len(self.mappings)}")
        keys   = sum(1 for m in self.mappings if m.is_key)
        xforms = sum(1 for m in self.mappings if m.has_transform)
        print(f"   Keys   : {keys}")
        print(f"   With transformations: {xforms}")
        print()
        for tbl in self.tables:
            fields = self.for_table(tbl)
            tgt = fields[0].target_table if fields else "?"
            print(f"   {tbl} → {tgt}  ({len(fields)} fields)")
        if self.warnings:
            print(f"\n   ⚠️  Warnings:")
            for w in self.warnings:
                print(f"      • {w}")


def _resolve_header(headers: list[str]) -> dict[str, int]:
    """Map canonical field names to column indices, ignoring case/spaces."""
    h_lower = {h.strip().lower(): i for i, h in enumerate(headers)}
    resolved = {}
    for canonical, aliases in _COL_ALIASES.items():
        for alias in aliases:
            if alias.lower() in h_lower:
                resolved[canonical] = h_lower[alias.lower()]
                break
    return resolved


def _parse_bool(val: str) -> bool:
    return str(val).strip().lower() in ("yes", "true", "1", "y", "x", "✓", "key")


def load_mapping(path: str) -> MappingDocument:
    """
    Load a Functional/Technical mapping document from Excel or CSV.

    Args:
        path: Path to .xlsx or .csv mapping file.

    Returns:
        MappingDocument with all field mappings parsed.

    Example:
        from reconpilot.mapping import load_mapping
        doc = load_mapping("mapping/CRM_Field_Mapping.xlsx")
        doc.summary()
    """
    p = Path(path)
    warnings = []
    rows = []

    if p.suffix.lower() in (".xlsx", ".xls"):
        try:
            import openpyxl
        except ImportError:
            raise ImportError("pip install pipe-recon[excel] to read Excel mapping files")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        # Use first sheet by default, or sheet named "Mapping" / "Field Mapping"
        sheet = None
        for name in ["Mapping", "Field Mapping", "FieldMapping", "Technical Mapping"]:
            if name in wb.sheetnames:
                sheet = wb[name]
                break
        if sheet is None:
            sheet = wb.active
        data = list(sheet.values)
        if not data:
            raise ValueError("Empty sheet in mapping document")
        headers = [str(h or "").strip() for h in data[0]]
        rows = [[str(c or "").strip() for c in row] for row in data[1:] if any(c for c in row)]

    elif p.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            data = list(reader)
        if not data:
            raise ValueError("Empty CSV mapping file")
        headers = [h.strip() for h in data[0]]
        rows = [[c.strip() for c in row] for row in data[1:] if any(c.strip() for c in row)]

    else:
        raise ValueError(f"Unsupported mapping file format: {p.suffix}. Use .xlsx or .csv")

    col = _resolve_header(headers)

    if "source_table" not in col or "source_column" not in col:
        raise ValueError(
            f"Could not find source_table/source_column in mapping doc. "
            f"Headers found: {headers}"
        )
    if "target_table" not in col or "target_column" not in col:
        raise ValueError(
            f"Could not find target_table/target_column in mapping doc. "
            f"Headers found: {headers}"
        )

    mappings = []
    for i, row in enumerate(rows, 2):
        def get(key, default=""):
            idx = col.get(key)
            if idx is None or idx >= len(row):
                return default
            return row[idx].strip()

        src_tbl = get("source_table")
        src_col = get("source_column")
        tgt_tbl = get("target_table")
        tgt_col = get("target_column")

        if not src_tbl or not src_col or not tgt_tbl or not tgt_col:
            warnings.append(f"Row {i}: skipped (missing required fields)")
            continue

        m = FieldMapping(
            source_table=src_tbl,
            source_column=src_col,
            target_table=tgt_tbl,
            target_column=tgt_col,
            transformation=get("transformation", "direct"),
            is_key=_parse_bool(get("is_key", "no")),
            nullable=_parse_bool(get("nullable", "yes")),
            data_type=get("data_type", ""),
            notes=get("notes", ""),
        )
        mappings.append(m)

    if not mappings:
        raise ValueError("No valid field mappings found in document")

    return MappingDocument(source_path=str(path), mappings=mappings, warnings=warnings)


def create_sample_mapping(output_path: str = "sample_mapping.csv"):
    """Write a sample CSV mapping file you can fill in and use immediately."""
    rows = [
        ["source_table", "source_column", "target_table", "target_column",
         "transformation", "is_key", "nullable", "data_type", "notes"],
        ["CUSTOMER", "CUST_ID",   "customers", "customer_id",  "direct", "yes", "no",  "INTEGER", "Primary key"],
        ["CUSTOMER", "FIRST_NM",  "customers", "first_name",   "trim",   "no",  "no",  "VARCHAR", ""],
        ["CUSTOMER", "LAST_NM",   "customers", "last_name",    "trim",   "no",  "no",  "VARCHAR", ""],
        ["CUSTOMER", "EMAIL_ADDR","customers", "email",        "lower",  "no",  "yes", "VARCHAR", "Normalize to lowercase"],
        ["CUSTOMER", "BIRTH_DT",  "customers", "birth_date",   "date_format: YYYYMMDD→YYYY-MM-DD", "no", "yes", "DATE", ""],
        ["CUSTOMER", "STAT_CD",   "customers", "status",       "lookup: A=Active,I=Inactive", "no", "no", "VARCHAR", ""],
        ["ORDER",    "ORD_ID",    "orders",    "order_id",     "direct", "yes", "no",  "INTEGER", ""],
        ["ORDER",    "CUST_ID",   "orders",    "customer_id",  "direct", "no",  "no",  "INTEGER", "FK to customers"],
        ["ORDER",    "ORD_DT",    "orders",    "order_date",   "date_format: YYYYMMDD→YYYY-MM-DD", "no", "no", "DATE", ""],
        ["ORDER",    "TOT_AMT",   "orders",    "total_amount", "direct", "no",  "no",  "DECIMAL", ""],
    ]
    with open(output_path, "w", newline="") as f:
        csv.writer(f).writerows(rows)
    print(f"✅ Sample mapping created: {output_path}")
    print("   Fill this in with your actual field mappings and run:")
    print("   pipe-recon reconcile mapping.csv --source <dsn> --target <dsn>")
